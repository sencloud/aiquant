import re, sys
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

src, dst = sys.argv[1], sys.argv[2]
doc = Document()
st = doc.styles['Normal']; st.font.name = 'Noto Sans CJK SC'; st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'PingFang SC')

def runs(p, text):
    for part in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text):
        if not part: continue
        if part.startswith('**'): p.add_run(part[2:-2]).bold = True
        elif part.startswith('`'):
            r = p.add_run(part[1:-1]); r.font.name = 'Menlo'
        else: p.add_run(part)

lines = open(src, encoding='utf-8').read().split('\n'); i = 0
while i < len(lines):
    l = lines[i]
    if l.startswith('```'):
        j = i + 1; buf = []
        while j < len(lines) and not lines[j].startswith('```'): buf.append(lines[j]); j += 1
        t = doc.add_table(rows=1, cols=1); t.style = 'Table Grid'
        c = t.cell(0, 0); c.text = ''
        for k, b in enumerate(buf):
            p = c.paragraphs[0] if k == 0 else c.add_paragraph()
            p.add_run(b)
        doc.add_paragraph(); i = j + 1; continue
    if l.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].startswith('|'):
            cells = [x.strip() for x in lines[i].strip('|').split('|')]
            if not all(re.fullmatch(r'-+', x) for x in cells): rows.append(cells)
            i += 1
        t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = 'Table Grid'
        for r, row in enumerate(rows):
            for c, val in enumerate(row):
                cell = t.cell(r, c); cell.text = ''
                runs(cell.paragraphs[0], val)
                if r == 0:
                    for run in cell.paragraphs[0].runs: run.bold = True
        doc.add_paragraph(); continue
    m = re.match(r'(#+) (.*)', l)
    if m: doc.add_heading(m.group(2), level=min(len(m.group(1)), 3)); i += 1; continue
    if l.startswith('> '): runs(doc.add_paragraph(style='Intense Quote'), l[2:].rstrip()); i += 1; continue
    if re.match(r'\s*[-·] ', l): runs(doc.add_paragraph(style='List Bullet'), re.sub(r'^\s*[-·] ', '', l)); i += 1; continue
    if re.match(r'\d+\. ', l): runs(doc.add_paragraph(style='List Number'), re.sub(r'^\d+\. ', '', l)); i += 1; continue
    if l.strip(): runs(doc.add_paragraph(), l)
    i += 1
doc.save(dst)
